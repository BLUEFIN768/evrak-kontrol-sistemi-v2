import openpyxl
from dataclasses import dataclass
from typing import Optional, List, Dict
from datetime import datetime
import re

@dataclass
class Evidence:
    code: str
    document_no: str
    issue_date: Optional[datetime] = None
    expiry_date: Optional[datetime] = None
    verified: bool = False
    raw_text: str = ""

@dataclass
class Result:
    code: str
    status: str
    message: str
    evidence_used: Optional[Evidence] = None

def load_rules_from_excel(filepath: str, rank: str, is_tanker: bool = False) -> Dict[str, str]:
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb["DOC_MATRIX"] if "DOC_MATRIX" in wb.sheetnames else wb.active
        
        rank_col = 6
        rank_clean = rank.strip().lower()
        for col in range(6, ws.max_column + 1):
            val = ws.cell(row=9, column=col).value
            if val and str(val).strip().lower() in rank_clean:
                rank_col = col
                break
                
        rules = {}
        for row in range(14, ws.max_row + 1):
            doc_name = ws.cell(row=row, column=2).value
            rule_val = ws.cell(row=row, column=rank_col).value
            
            if doc_name and rule_val:
                rule_str = str(rule_val).strip().upper()
                if rule_str in ["X", "D"]:
                    doc_clean = str(doc_name).strip().upper()
                    matched_code = None
                    
                    if "PASSPORT" in doc_clean or "PASAPORT" in doc_clean: matched_code = "PASSPORT"
                    elif "SEAMAN" in doc_clean or "CÜZDAN" in doc_clean: matched_code = "SEAMAN_BOOK"
                    elif "MEDICAL" in doc_clean or "SAĞLIK" in doc_clean: matched_code = "MEDICAL_CERT"
                    elif "YELLOW" in doc_clean or "FEVER" in doc_clean: matched_code = "YELLOW_FEVER"
                    elif "STCW" in doc_clean or "BASIC" in doc_clean or "SAFETY" in doc_clean: matched_code = "BASIC_SAFETY"
                    elif "GMDSS" in doc_clean or "RADIO" in doc_clean: matched_code = "GMDSS"
                    elif "COMPETENCY" in doc_clean or "YETERLİK" in doc_clean or "EHLİYET" in doc_clean: matched_code = "COMPETENCY"
                    elif "CRIMINAL" in doc_clean or "SİCİL" in doc_clean: matched_code = "CRIMINAL_RECORD"
                    elif "ECDIS" in doc_clean: matched_code = "ECDIS_GENERIC"
                    
                    if not matched_code:
                        matched_code = doc_clean.replace(" ", "_")
                        
                    rules[matched_code] = rule_str
                    
        return rules
    except Exception as e:
        print(f"Excel rule read error: {e}")
        return {}

def _check_dates(evidence: Evidence, rule_type: str) -> str:
    if not evidence.issue_date and evidence.code != "YELLOW_FEVER":
        return "REVIEW"
    if evidence.code == "YELLOW_FEVER":
        return "OK"
    if evidence.expiry_date and evidence.expiry_date < datetime.now():
        return "EXPIRED"
    return "OK"

def evaluate(evidences: List[Evidence], rank: str, excel_path: str, is_tanker: bool = False) -> List[Result]:
    rules = load_rules_from_excel(excel_path, rank, is_tanker)
    results = []
    
    if not rules:
        return [Result("ERROR", "MISSING", "Kurallar okunamadı veya rütbe bulunamadı.")]
    
    for req_code, rule_type in rules.items():
        if "TANKER" in req_code and not is_tanker:
            continue
            
        matched_evidence = None
        for ev in evidences:
            if ev.code == req_code or req_code in ev.code or ev.code in req_code:
                matched_evidence = ev
                break
                
        if not matched_evidence:
            if rule_type == 'D':
                results.append(Result(code=req_code, status="OK", message="Atamaya Bağlı - Belge Yok", evidence_used=None))
            else:
                results.append(Result(code=req_code, status="MISSING", message="Zorunlu Belge Eksik", evidence_used=None))
            continue
            
        date_status = _check_dates(matched_evidence, rule_type)
        if date_status != "OK":
            msg = "Belge Süresi Dolmuş" if date_status == "EXPIRED" else "Tarih Okunamadı veya Eksik"
            results.append(Result(code=req_code, status=date_status, message=msg, evidence_used=matched_evidence))
            continue
            
        if not matched_evidence.verified:
            results.append(Result(code=req_code, status="REVIEW", message="Orijinalden Doğrulanmalı (Kanıtsız Yes Yok)", evidence_used=matched_evidence))
            continue
            
        results.append(Result(code=req_code, status="OK", message="Geçerli", evidence_used=matched_evidence))
        
    return results
