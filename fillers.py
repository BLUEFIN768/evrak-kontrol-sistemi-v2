import openpyxl
from docx import Document
import io
from typing import List, Dict
from crew_core import Result

def safe_write_excel(ws, cell_coord, val):
    try:
        cell = ws[cell_coord]
        if type(cell).__name__ == 'MergedCell':
            for r in ws.merged_cells.ranges:
                if cell.coordinate in r:
                    ws.cell(row=r.min_row, column=r.min_col).value = val
                    return
        ws[cell_coord].value = val
    except Exception:
        pass

def fill_excel(template_path: str, output_stream, results: List[Result], personel_info: dict):
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    
    safe_write_excel(ws, "C3", personel_info.get("name", ""))
    safe_write_excel(ws, "C4", personel_info.get("rank", ""))
    safe_write_excel(ws, "G3", personel_info.get("vessel", ""))
    safe_write_excel(ws, "G4", personel_info.get("date", ""))
    
    row_idx = 7
    for res in results:
        ws.cell(row=row_idx, column=2).value = res.code
        ws.cell(row=row_idx, column=5).value = res.status
        ws.cell(row=row_idx, column=7).value = res.message
        row_idx += 1
        
    wb.save(output_stream)

def fill_word(template_path: str, output_stream, results: List[Result], personel_info: dict, operator_name: str):
    doc = Document(template_path)
    replacements = {
        "{{NAME}}": personel_info.get("name", ""),
        "{{RANK}}": personel_info.get("rank", ""),
        "{{VESSEL}}": personel_info.get("vessel", ""),
        "{{DATE}}": personel_info.get("date", ""),
        "{{OPERATOR}}": operator_name
    }
    
    for p in doc.paragraphs:
        for k, v in replacements.items():
            if k in p.text:
                p.text = p.text.replace(k, v)
                
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for k, v in replacements.items():
                    if k in cell.text:
                        cell.text = cell.text.replace(k, v)
                        
    doc.save(output_stream)
